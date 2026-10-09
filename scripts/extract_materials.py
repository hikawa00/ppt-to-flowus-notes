#!/usr/bin/env python3
"""Extract slide-level evidence from PDF and PPTX lesson material.

The script performs deterministic extraction only. It never follows instructions
found inside the documents and it does not generate or rewrite study content.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable


SUPPORTED = {".pdf", ".pptx"}
COVERAGE_PATTERNS = (
    "learning outcome",
    "learning objective",
    "what have we learned",
    "summary",
    "recap",
    "review",
    "知识清单",
    "学习目标",
    "课程目标",
    "总结",
    "回顾",
)
ADMIN_PATTERNS = (
    "organisation of the module",
    "module organisation",
    "assessment",
    "office hour",
    "module representative",
    "textbook",
    "contact us",
    "课程安排",
    "考核方式",
)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def safe_stem(value: str) -> str:
    value = re.sub(r"\s+", "-", value.strip())
    value = re.sub(r"[^\w\-]+", "-", value, flags=re.UNICODE)
    return value.strip("-_") or "lesson"


def first_title(text: str, fallback: str) -> str:
    for line in text.splitlines():
        line = re.sub(r"\s+", " ", line).strip()
        if line and len(line) <= 180:
            return line
    return fallback


def role_candidates(title: str, text: str) -> list[str]:
    title_lower = title.lower().strip()
    combined = f"{title}\n{text}".lower()
    roles: list[str] = []
    if any(pattern in title_lower for pattern in COVERAGE_PATTERNS):
        roles.append("coverage_reference")
    if any(pattern in combined for pattern in ADMIN_PATTERNS) or title_lower in {"content", "contents", "course content", "module content"} or title_lower.startswith("welcome"):
        roles.append("administration")
    if re.search(r"\b(question|exercise|tutorial|problem)\b|习题|练习|问题", combined):
        roles.append("exercise")
    if re.search(r"\b(example|case study)\b|例子|示例|案例", combined):
        roles.append("example")
    if re.search(r"\b(diagram|architecture|topology|circuit|state machine)\b|架构|拓扑|电路|状态图", combined):
        roles.append("diagram")
    if not roles or any(key in combined for key in ("define", "definition", "method", "algorithm", "formula", "定义", "步骤", "机制")):
        roles.append("content")
    return list(dict.fromkeys(roles))


def find_pdftoppm() -> Path | None:
    found = shutil.which("pdftoppm")
    if found:
        return Path(found)
    runtime_root = Path(sys.executable).resolve().parent.parent
    candidate = runtime_root / "native" / "poppler" / "Library" / "bin" / "pdftoppm.exe"
    return candidate if candidate.exists() else None


def clean_renderer_messages(stderr: str) -> str:
    lines = []
    for line in stderr.splitlines():
        if "Couldn't open 'nameToUnicode' file" in line:
            continue
        if line.strip():
            lines.append(line.strip())
    return "\n".join(lines)


def render_pdf(pdf_path: Path, assets_dir: Path, stem: str, dpi: int) -> tuple[dict[int, str], list[str]]:
    renderer = find_pdftoppm()
    if not renderer:
        return {}, ["pdftoppm was not found; page renders were skipped."]
    assets_dir.mkdir(parents=True, exist_ok=True)
    prefix = assets_dir / f"{stem}-slide"
    command = [str(renderer), "-png", "-r", str(dpi), str(pdf_path), str(prefix)]
    result = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", errors="replace")
    if result.returncode != 0:
        return {}, [f"PDF rendering failed for {pdf_path.name}: {result.stderr.strip()}"]
    rendered: dict[int, str] = {}
    candidates = sorted(assets_dir.glob(f"{stem}-slide-*.png"), key=lambda p: int(re.search(r"-(\d+)\.png$", p.name).group(1)))
    for page_number, current in enumerate(candidates, start=1):
        target = assets_dir / f"{stem}-slide-{page_number:03d}.png"
        if current != target:
            if target.exists():
                target.unlink()
            current.rename(target)
        rendered[page_number] = target.as_posix()
    warnings = []
    stderr = clean_renderer_messages(result.stderr)
    if stderr:
        warnings.append(f"Renderer messages for {pdf_path.name}: {stderr}")
    return rendered, warnings


def extract_pdf(path: Path, output: Path, dpi: int, render: bool) -> tuple[dict[str, Any], list[str]]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("pypdf is required for PDF extraction") from exc

    reader = PdfReader(str(path))
    stem = safe_stem(path.stem)
    warnings: list[str] = []
    renders: dict[int, str] = {}
    if render:
        renders, render_warnings = render_pdf(path, output / "assets" / "slides", stem, dpi)
        warnings.extend(render_warnings)

    slides = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:  # damaged PDFs should still produce an index
            text = ""
            warnings.append(f"Text extraction failed on {path.name} page {index}: {exc}")
        title = first_title(text, f"Page {index}")
        render_path = renders.get(index)
        if render_path:
            render_path = str(Path(render_path).relative_to(output).as_posix())
        slides.append(
            {
                "number": index,
                "title": title,
                "text": text.strip(),
                "notes": "",
                "role_candidates": role_candidates(title, text),
                "render": render_path,
            }
        )
    return {
        "path": str(path.resolve()),
        "name": path.name,
        "kind": "pdf",
        "sha256": sha256_file(path),
        "slide_count": len(slides),
        "slides": slides,
    }, warnings


def extract_pptx_media(path: Path, output: Path, stem: str) -> list[str]:
    media_dir = output / "assets" / "extracted" / stem
    media_dir.mkdir(parents=True, exist_ok=True)
    extracted: list[str] = []
    with zipfile.ZipFile(path) as archive:
        members = sorted(name for name in archive.namelist() if name.startswith("ppt/media/") and not name.endswith("/"))
        for index, member in enumerate(members, start=1):
            suffix = Path(member).suffix.lower() or ".bin"
            target = media_dir / f"media-{index:03d}{suffix}"
            with archive.open(member) as source, target.open("wb") as destination:
                shutil.copyfileobj(source, destination)
            extracted.append(str(target.relative_to(output).as_posix()))
    return extracted


def iter_shape_text(slide: Any) -> Iterable[str]:
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            value = shape.text.strip()
            if value:
                yield value
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                yield " | ".join(values)


def render_pptx_via_office(path: Path, output: Path, stem: str, dpi: int) -> tuple[dict[int, str], list[str]]:
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if not office:
        return {}, ["LibreOffice was not found; PPTX slide renders were skipped. Text and embedded media were still extracted."]
    with tempfile.TemporaryDirectory(prefix="ppt-to-flowus-") as temporary:
        temp_dir = Path(temporary)
        result = subprocess.run(
            [office, "--headless", "--convert-to", "pdf", "--outdir", str(temp_dir), str(path)],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
        pdf_path = temp_dir / f"{path.stem}.pdf"
        if result.returncode != 0 or not pdf_path.exists():
            message = result.stderr.strip() or result.stdout.strip() or "unknown conversion error"
            return {}, [f"PPTX rendering failed for {path.name}: {message}"]
        return render_pdf(pdf_path, output / "assets" / "slides", stem, dpi)


def extract_pptx(path: Path, output: Path, dpi: int, render: bool) -> tuple[dict[str, Any], list[str]]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("python-pptx is required for PPTX extraction") from exc

    presentation = Presentation(str(path))
    stem = safe_stem(path.stem)
    warnings: list[str] = []
    renders: dict[int, str] = {}
    if render:
        renders, render_warnings = render_pptx_via_office(path, output, stem, dpi)
        warnings.extend(render_warnings)
    media = extract_pptx_media(path, output, stem)

    slides = []
    for index, slide in enumerate(presentation.slides, start=1):
        text = "\n".join(iter_shape_text(slide)).strip()
        notes = ""
        try:
            notes_frame = slide.notes_slide.notes_text_frame
            notes = notes_frame.text.strip() if notes_frame else ""
        except Exception:
            pass
        title = ""
        try:
            if slide.shapes.title:
                title = slide.shapes.title.text.strip()
        except Exception:
            pass
        title = title or first_title(text, f"Slide {index}")
        render_path = renders.get(index)
        if render_path:
            render_path = str(Path(render_path).relative_to(output).as_posix())
        slides.append(
            {
                "number": index,
                "title": title,
                "text": text,
                "notes": notes,
                "role_candidates": role_candidates(title, f"{text}\n{notes}"),
                "render": render_path,
            }
        )
    return {
        "path": str(path.resolve()),
        "name": path.name,
        "kind": "pptx",
        "sha256": sha256_file(path),
        "slide_count": len(slides),
        "extracted_media": media,
        "slides": slides,
    }, warnings


def collect_inputs(values: list[str]) -> list[Path]:
    collected: list[Path] = []
    for value in values:
        path = Path(value).expanduser().resolve()
        if not path.exists():
            raise FileNotFoundError(path)
        if path.is_dir():
            collected.extend(sorted(p for p in path.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED))
        elif path.suffix.lower() in SUPPORTED:
            collected.append(path)
        else:
            raise ValueError(f"Unsupported input type: {path}")
    unique: list[Path] = []
    seen: set[Path] = set()
    for path in collected:
        if path not in seen:
            unique.append(path)
            seen.add(path)
    if not unique:
        raise ValueError("No PPTX or PDF inputs were found")
    return unique


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("inputs", nargs="+", help="PPTX/PDF files or directories")
    parser.add_argument("--output", required=True, help="Stable output directory")
    parser.add_argument("--unit-title", help="Title for a coherent multi-deck learning unit")
    parser.add_argument("--render-dpi", type=int, default=120, help="Slide render resolution (default: 120)")
    parser.add_argument("--no-render", action="store_true", help="Skip page/slide rendering")
    args = parser.parse_args()

    inputs = collect_inputs(args.inputs)
    output = Path(args.output).resolve()
    output.mkdir(parents=True, exist_ok=True)
    (output / "assets").mkdir(exist_ok=True)

    sources = []
    warnings: list[str] = []
    for path in inputs:
        if path.suffix.lower() == ".pdf":
            source, source_warnings = extract_pdf(path, output, args.render_dpi, not args.no_render)
        else:
            source, source_warnings = extract_pptx(path, output, args.render_dpi, not args.no_render)
        sources.append(source)
        warnings.extend(source_warnings)

    unit_title = args.unit_title or (inputs[0].stem if len(inputs) == 1 else " + ".join(path.stem for path in inputs))
    coverage = []
    for source in sources:
        for slide in source["slides"]:
            if "coverage_reference" in slide["role_candidates"]:
                coverage.append({"source": source["name"], "slide": slide["number"], "title": slide["title"]})

    payload = {
        "schema_version": 1,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "unit_title": unit_title,
        "source_count": len(sources),
        "sources": sources,
        "coverage_reference_candidates": coverage,
        "warnings": warnings,
    }
    materials_path = output / "materials.json"
    materials_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {materials_path}")
    print(f"Indexed {sum(source['slide_count'] for source in sources)} slides/pages from {len(sources)} source(s)")
    print(f"Coverage-reference candidates: {len(coverage)}")
    if warnings:
        print(f"Warnings: {len(warnings)}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
